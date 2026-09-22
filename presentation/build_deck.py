"""Builds the Group 7 presentation.

    python presentation/build_deck.py

Reads the figures in presentation/figures and results/plots, and writes
presentation/Group7_Assignment1_Incremental_Learning.pptx

The deck is written to be followed by someone who has not seen the project:
each slide title states the conclusion, and every term is explained before it
is used.
"""

import os

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_TICK_MARK
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
FIG = os.path.join(HERE, 'figures')
PLOTS = os.path.join(ROOT, 'results', 'plots')

DEEP = RGBColor(0x0B, 0x2E, 0x33)
TEAL = RGBColor(0x0E, 0x7C, 0x7B)
SEA = RGBColor(0x17, 0xBE, 0xBB)
MINT = RGBColor(0xD6, 0xF5, 0xF3)
PALE = RGBColor(0xEA, 0xF7, 0xF6)
CORAL = RGBColor(0xC9, 0x50, 0x2F)
CORAL_L = RGBColor(0xF2, 0x82, 0x5B)
INK = RGBColor(0x1C, 0x37, 0x38)
MUTED = RGBColor(0x55, 0x75, 0x75)
DIM = RGBColor(0x8F, 0xB5, 0xB5)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GRID = RGBColor(0xDC, 0xE9, 0xE8)

HEAD, BODY = 'Cambria', 'Calibri'

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]


# ------------------------------------------------------------------ helpers

def slide(dark=False):
    s = prs.slides.add_slide(BLANK)
    if dark:
        fill = s.background.fill
        fill.solid()
        fill.fore_color.rgb = DEEP
    return s


def text(s, x, y, w, h, runs, size=18, color=INK, font=BODY, bold=False,
         italic=False, align=PP_ALIGN.LEFT, spacing=1.0, space_after=0,
         bullets=False, anchor=MSO_ANCHOR.TOP):
    box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0

    for i, item in enumerate([runs] if isinstance(runs, str) else runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        if space_after:
            p.space_after = Pt(space_after)
        r = p.add_run()
        r.text = ('•  ' + item) if bullets else item
        f = r.font
        f.name, f.size, f.bold, f.italic = font, Pt(size), bold, italic
        f.color.rgb = color
    return box


def card(s, x, y, w, h, fill=PALE):
    shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                             Inches(x), Inches(y), Inches(w), Inches(h))
    shp.adjustments[0] = 0.06
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.color.rgb = fill
    shp.line.width = Pt(1)
    shp.shadow.inherit = False
    shp.text_frame.text = ''
    return shp


def stat(s, x, y, w, h, value, label, fill=PALE, value_color=TEAL,
         label_color=MUTED, value_size=40, label_size=14):
    card(s, x, y, w, h, fill=fill)
    text(s, x + 0.08, y + h * 0.13, w - 0.16, h * 0.5, value, size=value_size,
         color=value_color, font=HEAD, bold=True, align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + 0.12, y + h * 0.66, w - 0.24, h * 0.3, label, size=label_size,
         color=label_color, align=PP_ALIGN.CENTER)


def heading(s, title, kicker=None):
    if kicker:
        text(s, 0.65, 0.36, 12.0, 0.34, kicker.upper(), size=14, color=CORAL,
             bold=True)
    text(s, 0.65, 0.68 if kicker else 0.5, 12.0, 0.95, title, size=38,
         color=DEEP, font=HEAD, bold=True)


def explainer(s, x, y, w, body, label='WHAT THIS MEANS'):
    """A tinted box that explains a technique in plain words."""
    card(s, x, y, w, 1.55, fill=MINT)
    text(s, x + 0.35, y + 0.22, w - 0.7, 0.3, label, size=12, color=TEAL,
         bold=True)
    text(s, x + 0.35, y + 0.58, w - 0.7, 0.85, body, size=16, color=DEEP,
         spacing=1.15)


def divider(kicker, headline, sub=None):
    s = slide(dark=True)
    text(s, 1.1, 2.4, 11.1, 0.4, kicker.upper(), size=16, color=CORAL_L,
         bold=True)
    text(s, 1.1, 2.92, 11.1, 1.3, headline, size=46, color=WHITE, font=HEAD,
         bold=True)
    if sub:
        text(s, 1.1, 4.45, 11.1, 0.8, sub, size=20, color=DIM, spacing=1.15)
    return s


def numbered(s, x, y, w, n, head_text, detail):
    circ = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y),
                              Inches(0.72), Inches(0.72))
    circ.fill.solid()
    circ.fill.fore_color.rgb = TEAL
    circ.line.color.rgb = TEAL
    circ.shadow.inherit = False
    tf = circ.text_frame
    tf.text = str(n)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    f = p.runs[0].font
    f.name, f.size, f.bold, f.color.rgb = HEAD, Pt(23), True, WHITE

    text(s, x + 1.0, y - 0.05, w - 1.0, 0.42, head_text, size=21, color=DEEP,
         bold=True)
    text(s, x + 1.0, y + 0.44, w - 1.0, 0.9, detail, size=16, color=INK,
         spacing=1.15)


def picture(s, name, x, y, w, aspect, folder=PLOTS):
    s.shapes.add_picture(os.path.join(folder, name), Inches(x), Inches(y),
                         Inches(w), Inches(w / aspect))


def notes(s, txt):
    s.notes_slide.notes_text_frame.text = txt


# ------------------------------------------------------------------ 1. title

s = slide(dark=True)
text(s, 1.0, 1.9, 11.3, 1.05, 'Incremental Learning', size=54, color=WHITE,
     font=HEAD, bold=True)
text(s, 1.0, 2.88, 11.3, 1.05, 'on Large Data', size=54, color=SEA, font=HEAD,
     bold=True)
text(s, 1.0, 4.1, 11.3, 0.5,
     'Teaching a model to predict sales from a file too big to open',
     size=21, color=MINT)
text(s, 1.0, 5.35, 5.0, 0.4, 'GROUP 7', size=17, color=CORAL_L, bold=True)
text(s, 1.0, 5.78, 9.0, 0.4,
     'Ansley Smith   ·   Rebecca Drake   ·   Ximin Zeng', size=18,
     color=WHITE)
text(s, 1.0, 6.24, 9.0, 0.4,
     'BZAN 554  ·  Deep Learning  ·  Group Assignment 1', size=15,
     color=DIM)
notes(s, 'Our job was to predict how many units a product sells, using a file '
         'we were told to treat as too big to fit in the computer.')

# ------------------------------------------------------------- 2. the question

s = slide()
heading(s, 'The question: how many will this sell?', 'why anyone cares')
text(s, 0.65, 2.0, 7.3, 2.6,
     'An online retailer lists thousands of products. Each one sits on the '
     'site for a while, sells some number of units, then goes out of stock.',
     size=21, color=INK, spacing=1.25)
text(s, 0.65, 3.6, 7.3, 1.6,
     'If you could predict how many units a product will sell before you '
     'list it, you would know how much to stock and what to charge.',
     size=21, color=INK, spacing=1.25)
stat(s, 8.5, 2.0, 4.2, 2.3, 'quantity', 'the number we are predicting',
     fill=DEEP, value_color=SEA, label_color=DIM, value_size=36, label_size=15)
text(s, 8.5, 4.6, 4.2, 1.2,
     'Everything else in the data is a clue we are allowed to use.',
     size=17, color=CORAL, bold=True, spacing=1.2)
notes(s, 'Start with why this matters. A retailer wants to know, before '
         'listing a product, how many units it will move. Everything else in '
         'the file is a clue.')

# ------------------------------------------------------------------ 3. the data

s = slide()
heading(s, 'Each row is one product, for one spell in stock',
        'what the data looks like')

rows = [['sku', 'price', 'order', 'duration', 'category', 'quantity'],
        ['11', '1.45', '0', '1.50', '17', '5'],
        ['11', '1.45', '1', '2.47', '17', '15'],
        ['17', '0.66', '1', '1.85', '17', '11']]
tbl_shape = s.shapes.add_table(4, 6, Inches(0.65), Inches(1.9),
                               Inches(7.3), Inches(2.0))
tbl = tbl_shape.table
for c, width in enumerate([1.05, 1.15, 1.1, 1.35, 1.4, 1.25]):
    tbl.columns[c].width = Inches(width)
for r, row in enumerate(rows):
    tbl.rows[r].height = Inches(0.5)
    for c, val in enumerate(row):
        cell = tbl.cell(r, c)
        cell.text = val
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.fill.solid()
        if r == 0:
            cell.fill.fore_color.rgb = DEEP
        elif c == 5:
            cell.fill.fore_color.rgb = MINT
        else:
            cell.fill.fore_color.rgb = WHITE
        p = cell.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        f = p.runs[0].font
        f.name, f.size = BODY, Pt(15)
        f.bold = (r == 0 or c == 5)
        f.color.rgb = WHITE if r == 0 else (TEAL if c == 5 else INK)

text(s, 0.65, 4.2, 7.3, 0.5, 'The last column is what we predict. The other '
                             'five are what we predict it from.',
     size=17, color=CORAL, bold=True)

text(s, 8.5, 1.95, 4.3, 4.4, [
    'sku  —  which product it is',
    'price  —  what it costs',
    'order  —  how many times it has been listed before',
    'duration  —  how long it stayed on the site',
    'category  —  what kind of product',
], size=16, color=INK, spacing=1.2, space_after=14, bullets=True)
notes(s, 'Three real rows. A product can appear more than once, because a new '
         'row is created each time it comes back into stock. Duration is how '
         'long it stayed listed before selling out.')

# ------------------------------------------------------------------ 4. divider

s = divider('but there was a catch', 'We were not allowed to\nopen the file',
            'The assignment says to assume the data is bigger than the '
            'computer’s memory')
notes(s, 'This is the whole point of the assignment. The file we were given is '
         'only 18 MB, but we had to build as though it were far too big to '
         'load.')

# ------------------------------------------------------------- 5. the catch

s = slide()
heading(s, 'Normally you load everything. We could not.',
        'the constraint')
picture(s, 'contrast.png', 0.55, 1.85, 12.25, 2.660, folder=FIG)
text(s, 0.65, 6.35, 12.0, 0.6,
     'Reading a chunk at a time means the computer only ever holds a small '
     'piece — so the file could be a thousand times larger.',
     size=17, color=INK, spacing=1.15)
notes(s, 'The usual approach loads the whole file, then trains. That needs a '
         'machine big enough for the file. We read one chunk, learn from it, '
         'discard it, and read the next.')

# ------------------------------------------------------------- 6. the pipeline

s = slide()
heading(s, 'We read the file three times, a chunk at a time',
        'how it actually works')
picture(s, 'pipeline.png', 0.55, 1.8, 12.25, 2.556, folder=FIG)
text(s, 0.65, 6.6, 12.0, 0.6,
     'Reading from disk is cheap. Holding data in memory is what we cannot '
     'afford.', size=17, color=CORAL, bold=True)
notes(s, 'Pass one works out the average and spread of each column, which we '
         'need before we can put the numbers on a common scale. Pass two does '
         'the learning. Pass three sets aside rows to test on. Each pass holds '
         'only one chunk at a time.')

# ------------------------------------------------------------- 7. the model

s = slide()
heading(s, 'The model is a stack of pattern-finders', 'what we built')
picture(s, 'architecture.png', 0.55, 1.8, 12.25, 2.956, folder=FIG)
text(s, 0.65, 6.1, 12.0, 1.0,
     'Five facts go in. Each layer looks for patterns in what the previous '
     'layer found. One number comes out: how many units we expect to sell.',
     size=18, color=INK, spacing=1.2)
notes(s, 'A neural network is a chain of simple pattern-finders. Each layer '
         'takes what the last one noticed and looks for patterns in that. The '
         'assignment specified three hidden layers with sigmoid activation.')

# ------------------------------------------------------------- 8. decisions

s = slide()
heading(s, 'Three choices we had to make', 'and why')
numbered(s, 0.65, 1.95, 12.0, 1, 'We put the answer on the same scale as the clues',
         'Sales run from 1 to over 4,000. Numbers that large jam the model, so '
         'we shrank them for training and converted back afterwards.')
numbered(s, 0.65, 3.6, 12.0, 2, 'The last step had to be left unrestricted',
         'The building block we were told to use can only ever output a number '
         'between 0 and 1 — useless for predicting sales of 300.')
numbered(s, 0.65, 5.25, 12.0, 3, 'We had to mix up the rows before learning',
         'More on this next — it turned out to be the most important thing '
         'we did.')
notes(s, 'Three decisions we expect questions about. The third one is the '
         'interesting one and gets its own section.')

# ------------------------------------------------------------------ 9. divider

s = divider('then something went wrong', 'Our accuracy suddenly dropped',
            'Switching to reading the file in chunks made the model noticeably '
            'worse. Here is why.')
notes(s, 'Our first chunked version scored much worse than the version that '
         'loaded everything. It took a while to work out what was happening.')

# ------------------------------------------------------------- 10. sorted

s = slide()
heading(s, 'Every chunk looked almost the same', 'the problem we found')
text(s, 0.65, 1.95, 6.0, 3.6,
     'The rows are ordered by product. So a chunk read straight off disk is '
     'not a fair sample of the shop — it is a handful of nearly identical '
     'products in a row.', size=19, color=INK, spacing=1.25)
text(s, 0.65, 4.2, 6.0, 1.8,
     'The model kept adjusting to whatever narrow slice it happened to be '
     'looking at, then forgetting it when the next slice arrived.',
     size=19, color=INK, spacing=1.25)

stat(s, 7.2, 2.1, 5.5, 1.9, '0.9995',
     'correlation between row number and product id', fill=DEEP,
     value_color=SEA, label_color=DIM, value_size=52, label_size=15)
text(s, 7.2, 4.35, 5.5, 1.5,
     'In other words: the position of a row in the file tells you almost '
     'exactly which product it is.', size=18, color=CORAL, bold=True,
     spacing=1.2)
notes(s, 'Learning works by taking small steps based on small samples, and it '
         'assumes each sample is representative. Sorted data breaks that '
         'assumption completely.')

# ------------------------------------------------------------------ 11. the fix

s = slide()
heading(s, 'The fix: shuffle inside a small window', 'and it worked')
text(s, 0.65, 1.95, 5.3, 2.4,
     'We cannot shuffle a file we cannot open. So we hold a fixed number of '
     'rows, shuffle those, use them up, then refill.', size=18, color=INK,
     spacing=1.2)
stat(s, 0.65, 4.4, 5.3, 1.35, '0.44  →  0.53', 'accuracy, before and after',
     fill=MINT, value_size=30)
text(s, 0.65, 6.05, 5.3, 0.9,
     'The window is a fixed 200,000 rows — about 5 MB — no matter how '
     'big the file is.', size=15, color=MUTED, spacing=1.15)

chart_data = CategoryChartData()
chart_data.categories = ['no shuffle', '50,000', '100,000', '200,000', '400,000']
chart_data.add_series('Accuracy', (0.4445, 0.4478, 0.5000, 0.5312, 0.5406))
gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(6.3),
                        Inches(1.9), Inches(6.5), Inches(4.5), chart_data)
chart = gf.chart
chart.has_legend = False
chart.has_title = True
chart.chart_title.text_frame.text = 'Accuracy against size of the shuffle window'
f = chart.chart_title.text_frame.paragraphs[0].runs[0].font
f.name, f.size, f.bold, f.color.rgb = BODY, Pt(15), True, DEEP

plot = chart.plots[0]
plot.vary_by_categories = False
plot.gap_width = 60
plot.has_data_labels = True
dl = plot.data_labels
dl.number_format, dl.number_format_is_linked = '0.000', False
dl.position = XL_LABEL_POSITION.OUTSIDE_END
dl.font.name, dl.font.size, dl.font.bold = BODY, Pt(13), True
dl.font.color.rgb = INK
for idx, point in enumerate(plot.series[0].points):
    point.format.fill.solid()
    point.format.fill.fore_color.rgb = (CORAL if idx == 3 else
                                        MINT if idx == 4 else TEAL)
va = chart.value_axis
va.minimum_scale, va.maximum_scale = 0.40, 0.58
va.has_major_gridlines = True
va.major_gridlines.format.line.color.rgb = GRID
va.major_gridlines.format.line.width = Pt(0.75)
va.tick_labels.font.name, va.tick_labels.font.size = BODY, Pt(12)
va.tick_labels.font.color.rgb = MUTED
va.major_tick_mark = XL_TICK_MARK.NONE
ca = chart.category_axis
ca.has_major_gridlines = False
ca.tick_labels.font.name, ca.tick_labels.font.size = BODY, Pt(13)
ca.tick_labels.font.color.rgb = MUTED
ca.major_tick_mark = XL_TICK_MARK.NONE
text(s, 6.3, 6.55, 6.5, 0.45,
     'The last bar shuffles the whole file — shown only as a ceiling.',
     size=13, color=MUTED, italic=True)
notes(s, 'A bigger window mixes the rows more thoroughly and scores better, '
         'but costs more memory. At 200,000 rows we match what a full shuffle '
         'would give, for about five megabytes.')

# ----------------------------------------------------------------- 12. divider

s = divider('so', 'Did it work?',
            'Accuracy, what the model learned, and whether it stayed inside '
            'memory')
notes(s, 'Now the results the assignment asks for.')

# ------------------------------------------------------------- 13. accuracy

s = slide()
heading(s, 'It explains about half of why sales differ', 'accuracy')
card(s, 0.65, 1.95, 5.6, 3.0, fill=DEEP)
text(s, 0.85, 2.3, 5.2, 1.6, '0.53', size=100, color=SEA, font=HEAD, bold=True,
     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
text(s, 0.85, 4.05, 5.2, 0.5, 'R-SQUARED ON UNSEEN DATA', size=15, color=MINT,
     bold=True, align=PP_ALIGN.CENTER)

explainer(s, 6.6, 1.95, 6.1,
          'R-squared asks how much better we are than always guessing the '
          'average. 0 means no better. 1 means perfect.')
text(s, 6.6, 3.75, 6.1, 1.5,
     'At 0.53 the model explains roughly half of why some products sell more '
     'than others — and half is still unexplained.', size=17, color=INK,
     spacing=1.2)
stat(s, 6.6, 5.3, 2.95, 1.2, '0.523', 'on data it trained on', value_size=30,
     label_size=13.5)
stat(s, 9.75, 5.3, 2.95, 1.2, '0.528', 'on data it never saw', value_size=30,
     label_size=13.5)
text(s, 0.65, 5.2, 5.6, 1.1,
     'Those two scores are almost identical, which means the model is not '
     'simply memorising.', size=16, color=CORAL, bold=True, spacing=1.2)
notes(s, 'R-squared of 0.53. Train and test agree to within half a percent, so '
         'it is not memorising. Half the variation is still unexplained, which '
         'is honest for five columns and one pass over the data.')

# --------------------------------------------------------- 14. learning curve

s = slide()
heading(s, 'It learned fast, then levelled off', 'progress during training')
picture(s, 'learning_curve.png', 0.55, 1.85, 7.3, 1.705)
explainer(s, 8.2, 1.95, 4.55,
          'Each point is how wrong the model was, averaged over recent rows. '
          'Lower is better.')
text(s, 8.2, 3.75, 4.55, 2.0,
     'Error fell by about two-thirds in the first stretch, then flattened — '
     'the model had learned what one pass over the data could teach it.',
     size=17, color=INK, spacing=1.2)
stat(s, 8.2, 5.75, 4.55, 1.15, '2,831 → 989', 'average error, start to end',
     value_size=28, label_size=13.5)
notes(s, 'The x-axis is how many records the model has seen. Error drops '
         'steeply early, then flattens. A second pass would help a little, but '
         'true incremental learning sees each record once.')

# ------------------------------------------------------------- 15. importance

s = slide()
heading(s, 'How long it stayed listed matters most', 'which clues the model uses')
picture(s, 'variable_importance.png', 0.55, 1.95, 7.3, 1.733)
explainer(s, 8.2, 1.95, 4.55,
          'To test a clue, we scramble it and see how much worse the '
          'predictions get. Big drop means the model relied on it.')
text(s, 8.2, 3.8, 4.55, 2.4, [
    'duration is far and away the strongest',
    'price matters, but much less',
    'which product it is barely registers',
], size=17, color=DEEP, bold=True, spacing=1.15, space_after=14, bullets=True)
text(s, 8.2, 6.0, 4.55, 0.9,
     'How long something sits on the site tells you more than what it is.',
     size=16, color=CORAL, bold=True, spacing=1.15)
notes(s, 'Scrambling duration costs nearly all of our accuracy. Scrambling the '
         'product id costs essentially nothing, which means the model learned '
         'to ignore it.')

# ------------------------------------------------------------------ 16. pdp

s = slide()
heading(s, 'Longer listings sell far more', 'how each clue moves the prediction')
explainer(s, 0.65, 1.7, 12.0,
          'We pretend every product had the same value for one clue, predict, '
          'and average. Then we slide that value across its whole range.')
picture(s, 'partial_dependence.png', 0.55, 3.4, 12.25, 4.198)
text(s, 0.65, 6.45, 12.0, 0.6,
     'duration climbs steeply  ·  price drops off  ·  the rest are '
     'nearly flat', size=18, color=CORAL, bold=True)
notes(s, 'These five panels say which direction each clue pushes the '
         'prediction, and how hard. Duration spans nearly 390 units of '
         'predicted sales. Product id spans less than one.')

# ------------------------------------------------------------------ 17. RAM

s = slide()
heading(s, 'Memory never grew', 'did we stay inside the constraint')
picture(s, 'memory_usage.png', 0.55, 1.95, 7.2, 1.789)
explainer(s, 8.2, 1.95, 4.55,
          'If chunks were piling up in memory, this line would climb steadily. '
          'It does not.')
stat(s, 8.2, 3.8, 4.55, 1.25, '+10.7 MB', 'across all 400,000 records',
     value_size=34, label_size=13.5)
stat(s, 8.2, 5.25, 4.55, 1.25, '405 MB', 'peak, out of 14,171 available',
     value_size=34, label_size=13.5)
text(s, 0.65, 6.1, 7.4, 1.0,
     'The one-time step early on is the first chunk and the shuffle window '
     'being allocated. After that it is flat.', size=15, color=MUTED,
     italic=True, spacing=1.15)
notes(s, 'Two numbers here. About 40 MB is allocated once, at the start, for '
         'the first chunk, the shuffle window and TensorFlow itself. After '
         'that, across 400,000 records, memory moves by only 10.7 MB. That is '
         'the evidence that nothing is accumulating.')

# ------------------------------------------------------------------ 18. time

s = slide()
heading(s, 'The whole thing runs in under a minute', 'speed')
stat(s, 0.65, 2.1, 3.85, 2.2, '31 s', 'to learn from 400,000 records',
     value_size=56)
stat(s, 4.75, 2.1, 3.85, 2.2, '12,726', 'records per second', value_size=48)
stat(s, 8.85, 2.1, 3.85, 2.2, '40 s', 'start to finished plots', fill=DEEP,
     value_color=SEA, label_color=DIM, value_size=56)
text(s, 0.65, 4.8, 12.0, 2.0, [
    'One command produces the model, the scores and every plot in this deck',
    'At this rate, 100 million rows would take about two hours',
    'And it would use the same amount of memory it does now',
], size=19, spacing=1.15, space_after=16, bullets=True)
notes(s, 'Speed was never the constraint. Memory was. The point of the last '
         'line is that the approach scales: more data costs more time, not '
         'more memory.')

# ------------------------------------------------------------------ 19. next

s = slide(dark=True)
text(s, 1.0, 1.0, 11.3, 0.9, 'What we would do next', size=42, color=WHITE,
     font=HEAD, bold=True)
text(s, 1.0, 2.3, 11.3, 3.7, [
    'Treat product and category as labels, not numbers — right now the '
    'model is told product 5,000 is bigger than product 2,500',
    'Fix the random starting point so every run gives exactly the same answer',
    'Try a different building block in the middle layers; the one we were '
    'given learns slowly',
    'Measure what a second pass over the data would buy us',
], size=19, color=MINT, spacing=1.15, space_after=18, bullets=True)
text(s, 1.0, 6.5, 11.3, 0.45,
     'Group 7   ·   Ansley Smith, Rebecca Drake, Ximin Zeng', size=16,
     color=CORAL_L)
notes(s, 'The first one is the biggest missed opportunity, though the '
         'importance chart suggests the model already worked around it.')

# -------------------------------------------------------- 20. contributions

s = slide()
heading(s, 'Contributions', 'who did what')
for i, name in enumerate(['Ansley Smith', 'Rebecca Drake', 'Ximin Zeng']):
    x = 0.65 + i * 4.15
    card(s, x, 2.0, 3.85, 4.2)
    circ = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + 1.365), Inches(2.35),
                              Inches(1.12), Inches(1.12))
    circ.fill.solid()
    circ.fill.fore_color.rgb = TEAL
    circ.line.color.rgb = TEAL
    circ.shadow.inherit = False
    tf = circ.text_frame
    tf.text = ''.join(part[0] for part in name.split())
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    f = p.runs[0].font
    f.name, f.size, f.bold, f.color.rgb = HEAD, Pt(30), True, WHITE
    text(s, x + 0.15, 3.68, 3.55, 0.45, name, size=20, color=DEEP, bold=True,
         align=PP_ALIGN.CENTER)
    text(s, x + 0.2, 4.25, 3.45, 1.7, 'Fill in before submitting', size=15,
         color=MUTED, italic=True, align=PP_ALIGN.CENTER, spacing=1.2)
text(s, 0.65, 6.45, 12.0, 0.5,
     'Replace the three placeholders with each member’s actual '
     'contribution before submitting.', size=15, color=CORAL, bold=True)
notes(s, 'Fill this in as a team before submitting.')

# ------------------------------------------------------------------ write

OUT = os.path.join(HERE, 'Group7_Assignment1_Incremental_Learning.pptx')
prs.save(OUT)
print(f'wrote {OUT}')
print(f'{len(prs.slides._sldIdLst)} slides')
